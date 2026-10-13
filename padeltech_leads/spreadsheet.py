"""Build the Excel spreadsheet of every lead found so far.

It is regenerated from the system's memory on every run and attached to the
weekly email, so it is always complete and up to date. It is a read-only
export: edits made in a downloaded copy are not read back.
"""

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
BAND_FILL = {"HOT": "F8D7D3", "GOOD": "FCE5CD", "WATCH": "DDEBF7", "IGNORE": "EEEEEE"}
BAND_ORDER = {"HOT": 0, "GOOD": 1, "WATCH": 2, "IGNORE": 3}
STAGE_TEXT = {
    "submitted_pending": "Submitted – awaiting decision",
    "approved": "Permission granted",
    "condition_discharge": "Approved – conditions being discharged",
    "refused": "Refused",
    "withdrawn": "Withdrawn",
    "appeal": "At appeal",
    "unknown": "Unknown",
}

# (header, width, function lead -> value)
COLUMNS = [
    ("Band", 9, lambda l: l.get("band")),
    ("Score", 7, lambda l: l.get("score")),
    ("New this week", 9, None),  # filled in build()
    ("Lead", 44, lambda l: l.get("headline")),
    ("Council", 18, lambda l: l.get("council")),
    ("Application ref", 18, lambda l: l.get("reference")),
    ("Site address", 36, lambda l: l.get("address")),
    ("Postcode", 10, lambda l: l.get("postcode")),
    ("Courts", 8, lambda l: l.get("courts") if l.get("courts") is not None else "Unknown"),
    ("Indoor / outdoor", 13, lambda l: (l.get("indoor_outdoor") or "unknown").replace("_", " / ").capitalize()),
    ("Project type", 24, lambda l: (l.get("project_type") or "").replace("_", " ").capitalize()),
    ("Planning status", 24, lambda l: STAGE_TEXT.get(l.get("planning_stage"), "Unknown")),
    ("Applicant", 26, lambda l: l.get("applicant")),
    ("Planning agent", 26, lambda l: l.get("planning_agent")),
    ("Developer", 22, lambda l: l.get("developer")),
    ("Architect", 22, lambda l: l.get("architect")),
    ("Operator", 22, lambda l: l.get("operator")),
    ("Contacts found (public business info)", 46, lambda l: "\n".join(l.get("contacts") or []) or "Not researched"),
    ("Why this is a lead", 50, lambda l: l.get("why_lead")),
    ("Recommended action", 44, lambda l: l.get("recommended_action")),
    ("Update history", 40, lambda l: "\n".join(l.get("history") or [])),
    ("First found", 12, lambda l: date.fromisoformat(l["first_found"]) if l.get("first_found") else None),
    ("Last updated", 12, lambda l: date.fromisoformat(l["last_updated"]) if l.get("last_updated") else None),
    ("Planning application link", 30, lambda l: l.get("url")),
]


def build(path: str, leads: list[dict], today: date, week_start: date):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    body_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="1A5FB4", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")

    for c, (header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 32

    ordered = sorted(leads, key=lambda l: (BAND_ORDER.get(l.get("band"), 9), -(l.get("score") or 0),
                                           l.get("first_found") or ""))
    for r, lead in enumerate(ordered, start=2):
        new = (lead.get("first_found") or "") >= week_start.isoformat()
        fill = PatternFill("solid", fgColor=BAND_FILL.get(lead.get("band"), "FFFFFF"))
        for c, (header, _, get) in enumerate(COLUMNS, start=1):
            value = ("Yes" if new else "") if get is None else get(lead)
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = body_font, wrap
            if header in ("First found", "Last updated"):
                cell.number_format = "dd mmm yyyy"
            if header == "Band":
                cell.fill = fill
                cell.font = Font(name=FONT, size=10, bold=True)
            if header == "Planning application link" and value:
                cell.hyperlink = value
                cell.value = "Open application"
                cell.font = link_font

    last_row = max(2, len(ordered) + 1)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    _summary_sheet(wb, len(ordered), today)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _summary_sheet(wb: Workbook, count: int, today: date):
    ws = wb.create_sheet("Summary")
    bold = Font(name=FONT, bold=True)
    normal = Font(name=FONT)
    last = max(2, count + 1)
    ws["A1"] = "Padel lead database"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Updated automatically {today:%d %B %Y}. This file is regenerated every week, so "
    ws["A2"].font = normal
    ws["A3"] = "notes typed into a downloaded copy are not carried over."
    ws["A3"].font = normal

    rows = [
        ("Total leads", f"=COUNTA(Leads!A2:A{last})"),
        ("🔥 HOT (80–100)", f'=COUNTIF(Leads!A2:A{last},"HOT")'),
        ("🟠 GOOD (60–79)", f'=COUNTIF(Leads!A2:A{last},"GOOD")'),
        ("👀 WATCH (40–59)", f'=COUNTIF(Leads!A2:A{last},"WATCH")'),
        ("Dropped (e.g. refused/withdrawn)", f'=COUNTIF(Leads!A2:A{last},"IGNORE")'),
        ("New this week", f'=COUNTIF(Leads!C2:C{last},"Yes")'),
    ]
    for i, (label, formula) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label).font = bold
        ws.cell(row=i, column=2, value=formula).font = normal

    notes = [
        "How the score works (0–100):",
        "Padel relevance up to 35 · number of courts up to 15 · applicant type up to 15 · planning stage up to 15 ·",
        "construction imminent 5 · development size 5 · developer/agent identified 5 · likely to proceed 5.",
        "No padel connection is capped at 30; refused/withdrawn is capped at 35.",
        "",
        "Names marked '(AI inference)' in the email, and contacts marked 'inferred', are Claude's judgement –",
        "verify them before relying on them. Everything else comes from the planning record.",
        "Data: UK PlanIt (planit.org.uk) and council planning portals.",
    ]
    for i, text in enumerate(notes, start=13):
        ws.cell(row=i, column=1, value=text).font = bold if i == 13 else normal
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 12
